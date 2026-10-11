"""Regenerate the synthetic backend samples and their matching public copies."""

import shutil
from pathlib import Path

from docx import Document
from reportlab.lib.colors import HexColor
from reportlab.pdfgen.canvas import Canvas


def make_fixtures(destination):
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(
        str(destination / "project-proposal.pdf"), pagesize=(595, 842), invariant=True
    )
    for page in range(1, 4):
        canvas.setFillColor(HexColor("#182126"))
        canvas.setFont("Times-Bold", 27)
        canvas.drawString(
            62,
            753,
            "Q4 project proposal" if page == 1 else f"Project proposal - page {page}",
        )
        canvas.setStrokeColor(HexColor("#d8dfe5"))
        canvas.line(62, 731, 533, 731)
        canvas.setFont("Times-Roman", 15)
        canvas.drawString(62, 705, "A thoughtful plan for what comes next.")
        text = canvas.beginText(62, 655)
        text.setFont("Times-Roman", 12)
        text.setLeading(19)
        for line in [
            "This proposal outlines our plan for Q4, including key objectives,",
            "deliverables and the approach we will take to ensure a successful",
            "and sustainable outcome. It builds on recent progress and sets a",
            "clear path forward for the next phase of our work.",
        ]:
            text.textLine(line)
        canvas.drawText(text)
        for y, num, heading in [(535, "01", "Overview"), (345, "02", "Deliverables")]:
            canvas.setFillColor(HexColor("#276758"))
            canvas.setFont("Times-Roman", 20)
            canvas.drawString(62, y, num)
            canvas.setFillColor(HexColor("#182126"))
            canvas.setFont("Times-Bold", 20)
            canvas.drawString(107, y, heading)
            text = canvas.beginText(62, y - 35)
            text.setFont("Times-Roman", 12)
            text.setLeading(19)
            lines = (
                [
                    "In Q4, we will focus on delivering a more consistent experience",
                    "for our users. Our work will build on what we have learned,",
                    "address existing gaps, and support thoughtful growth.",
                    "",
                    "This plan is practical, focused and achievable with our resources.",
                ]
                if num == "01"
                else [
                    "1.   Refine and launch the core product improvements",
                    "2.   Enhance documentation and team support resources",
                    "3.   Gather and incorporate user feedback",
                    "4.   Prepare for a successful handover into next year",
                ]
            )
            for line in lines:
                text.textLine(line)
            canvas.drawText(text)
        canvas.setFont("Helvetica", 9)
        canvas.setFillColor(HexColor("#88939b"))
        canvas.drawString(62, 48, f"Papertrail sample document - Page {page} of 3")
        canvas.showPage()
    canvas.save()
    doc = Document()
    doc.add_heading("Quarterly project brief", 0)
    doc.add_paragraph("A Word document rendered inside Papertrail.")
    doc.add_heading("Overview", 1)
    doc.add_paragraph(
        "This proposal is ready for the reviewer to read and make a decision."
    )
    doc.add_heading("Deliverables", 1)
    for item in [
        "A clear submission process",
        "An accessible document viewer",
        "A recorded approval decision",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Milestone"
    table.rows[0].cells[1].text = "Status"
    row = table.add_row()
    row.cells[0].text = "Review"
    row.cells[1].text = "Ready"
    doc.add_page_break()
    doc.add_heading("Supporting details", 1)
    doc.add_paragraph("This second page verifies multi-page Word rendering.")
    doc.save(destination / "project-brief.docx")


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    samples = root / "workflow" / "samples"
    make_fixtures(samples)
    public = root / "frontend" / "public" / "samples"
    public.mkdir(parents=True, exist_ok=True)
    for filename in ("project-proposal.pdf", "project-brief.docx"):
        shutil.copyfile(samples / filename, public / filename)
