from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import IntegrityError
from docx import Document
from rest_framework.test import APIClient

from workflow.errors import StorageUnavailable
from workflow.models import Submission

pytestmark = pytest.mark.django_db


def test_login_session_logout_csrf(users):
    client = APIClient(enforce_csrf_checks=True)
    response = client.get("/api/session/")
    assert response.data["user"] is None
    assert (
        client.post(
            "/api/login/", {"username": "alex", "password": "Test-password-482!"}
        ).status_code
        == 403
    )
    token = response.data["csrfToken"]
    response = client.post(
        "/api/login/", {"username": "alex", "password": "bad"}, HTTP_X_CSRFTOKEN=token
    )
    assert response.status_code == 401
    response = client.post(
        "/api/login/",
        {"username": "alex", "password": "Test-password-482!"},
        HTTP_X_CSRFTOKEN=token,
    )
    assert response.status_code == 200
    assert response.data["user"]["role"] == "submitter"
    assert client.post("/api/logout/").status_code == 403
    assert (
        client.post(
            "/api/logout/", HTTP_X_CSRFTOKEN=response.data["csrfToken"]
        ).status_code
        == 200
    )
    assert client.get("/api/session/").data["user"] is None


def test_unauthenticated_access(client):
    assert client.get("/api/submissions/").status_code == 403


def test_creation_defaults_and_ownership(client, users, pdf, storage):
    client.force_authenticate(users[0])
    response = client.post(
        "/api/submissions/",
        {
            "title": "  Proposal  ",
            "file": pdf,
            "owner": users[1].pk,
            "status": "approved",
        },
        format="multipart",
    )
    assert response.status_code == 201
    assert response.data["title"] == "Proposal"
    assert response.data["status"] == "pending"
    assert response.data["owner"]["id"] == users[0].pk
    assert "blob_path" not in response.data
    assert storage[0].call_args.args[1].startswith(b"%PDF-")


def test_isolation_and_private_content(client, users, submission, storage, pdf_bytes):
    path = f"/api/submissions/{submission['id']}/"
    client.force_authenticate(users[1])
    assert client.get("/api/submissions/").data["results"] == []
    assert client.get(path).status_code == 404
    assert client.get(path + "file/").status_code == 404
    storage[1].assert_not_called()
    client.force_authenticate(users[2])
    assert len(client.get("/api/submissions/").data["results"]) == 1
    response = client.get(path + "file/")
    assert response.content == pdf_bytes
    assert response["Cache-Control"] == "private, no-store"
    assert response["Content-Type"] == "application/pdf"


@pytest.mark.parametrize("status", ["approved", "rejected"])
def test_decision_and_stale_conflict(client, users, submission, status):
    path = f"/api/submissions/{submission['id']}/decision/"
    client.force_authenticate(users[0])
    assert client.post(path, {"status": status}, format="json").status_code == 403
    client.force_authenticate(users[2])
    result = client.post(path, {"status": status}, format="json")
    assert result.status_code == 200
    assert result.data["status"] == status
    assert result.data["reviewed_by"]["id"] == users[2].pk
    assert result.data["reviewed_at"]
    other = "rejected" if status == "approved" else "approved"
    assert client.post(path, {"status": other}, format="json").status_code == 409
    assert client.post(path, {"status": status}, format="json").status_code == 409
    assert Submission.objects.get(pk=submission["id"]).status == status


def test_stale_objects_cannot_overwrite_decisions(users, submission):
    from workflow.errors import Conflict
    from workflow.services import decide_submission

    first = Submission.objects.get(pk=submission["id"])
    stale = Submission.objects.get(pk=submission["id"])
    decide_submission(users[2], first, "approved")
    with pytest.raises(Conflict):
        decide_submission(users[2], stale, "rejected")
    assert Submission.objects.get(pk=first.pk).status == "approved"


def test_invalid_decision_and_filter(client, users, submission):
    client.force_authenticate(users[2])
    assert (
        client.post(
            f"/api/submissions/{submission['id']}/decision/",
            {"status": "pending"},
            format="json",
        ).status_code
        == 400
    )
    assert client.get("/api/submissions/?status=invalid").status_code == 400
    assert len(client.get("/api/submissions/?status=pending").data["results"]) == 1
    assert client.get("/api/submissions/?status=approved").data["results"] == []


@pytest.mark.parametrize(
    "filename,content",
    [
        ("document.exe", b"content"),
        ("fake.pdf", b"not a PDF"),
        ("fake.docx", b"not a ZIP"),
        ("empty.pdf", b""),
        ("large.pdf", b"x" * 4_000_001),
    ],
)
def test_invalid_files(client, users, storage, filename, content):
    client.force_authenticate(users[0])
    response = client.post(
        "/api/submissions/",
        {"title": "Test", "file": SimpleUploadedFile(filename, content)},
        format="multipart",
    )
    assert response.status_code == 400
    storage[0].assert_not_called()


def test_title_and_file_required(client, users, storage, pdf):
    client.force_authenticate(users[0])
    assert (
        client.post(
            "/api/submissions/", {"title": "Missing file"}, format="multipart"
        ).status_code
        == 400
    )
    assert (
        client.post(
            "/api/submissions/", {"title": "  ", "file": pdf}, format="multipart"
        ).status_code
        == 400
    )
    storage[0].assert_not_called()


def test_multiple_files_rejected(client, users, pdf_bytes, storage):
    client.force_authenticate(users[0])
    files = [SimpleUploadedFile(f"{n}.pdf", pdf_bytes) for n in range(2)]
    assert (
        client.post(
            "/api/submissions/", {"title": "Two", "file": files}, format="multipart"
        ).status_code
        == 400
    )
    storage[0].assert_not_called()


def test_real_docx(client, users, storage):
    stream = BytesIO()
    document = Document()
    document.add_heading("Proposal", 0)
    document.add_paragraph("A real Word document.")
    document.save(stream)
    client.force_authenticate(users[0])
    response = client.post(
        "/api/submissions/",
        {
            "title": "Word proposal",
            "file": SimpleUploadedFile("proposal.docx", stream.getvalue()),
        },
        format="multipart",
    )
    assert response.status_code == 201
    assert response.data["content_type"].endswith("wordprocessingml.document")


def test_reviewer_cannot_submit(client, users, pdf, storage):
    client.force_authenticate(users[2])
    assert (
        client.post(
            "/api/submissions/",
            {"title": "Own review", "file": pdf},
            format="multipart",
        ).status_code
        == 403
    )
    storage[0].assert_not_called()


def test_upload_failure_creates_no_row(client, users, pdf, storage):
    storage[0].side_effect = StorageUnavailable()
    client.force_authenticate(users[0])
    assert (
        client.post(
            "/api/submissions/", {"title": "Proposal", "file": pdf}, format="multipart"
        ).status_code
        == 503
    )
    assert Submission.objects.count() == 0


def test_database_failure_cleans_up(client, users, pdf, storage):
    client.force_authenticate(users[0])
    with patch(
        "workflow.services.Submission.objects.create",
        side_effect=IntegrityError("failed"),
    ):
        assert (
            client.post(
                "/api/submissions/",
                {"title": "Proposal", "file": pdf},
                format="multipart",
            ).status_code
            == 503
        )
    storage[2].assert_called_once()
    assert Submission.objects.count() == 0


def test_cleanup_failure_is_logged(client, users, pdf, storage, caplog):
    storage[2].side_effect = RuntimeError("offline")
    client.force_authenticate(users[0])
    with patch(
        "workflow.services.Submission.objects.create",
        side_effect=IntegrityError("failed"),
    ):
        response = client.post(
            "/api/submissions/", {"title": "Proposal", "file": pdf}, format="multipart"
        )
    assert response.status_code == 503
    assert "Orphan document cleanup failed" in caplog.text


def test_preview_storage_failure(client, users, submission, storage):
    storage[1].side_effect = StorageUnavailable()
    assert client.get(submission["file_url"]).status_code == 503


def test_seeding_is_idempotent(monkeypatch):
    for name in ("SEED_ALEX_PASSWORD", "SEED_JAMIE_PASSWORD", "SEED_REVIEWER_PASSWORD"):
        monkeypatch.setenv(name, "Unique-test-password-924!")
    call_command("seed_users")
    call_command("seed_users")
    assert get_user_model().objects.count() == 3
    assert get_user_model().objects.filter(groups__name="reviewer").count() == 1


def test_missing_seed_credentials_is_atomic(monkeypatch):
    for name in ("SEED_ALEX_PASSWORD", "SEED_JAMIE_PASSWORD", "SEED_REVIEWER_PASSWORD"):
        monkeypatch.delenv(name, raising=False)
    with pytest.raises(CommandError):
        call_command("seed_users")
    assert get_user_model().objects.count() == 0


@pytest.mark.parametrize(
    "kind,expected", [("hyperlink", 201), ("image", 400), ("embedded", 400)]
)
def test_docx_external_resources_and_embedded_files(
    client, users, storage, kind, expected
):
    stream = BytesIO()
    document = Document()
    document.add_paragraph("Document with a relationship")
    document.save(stream)
    with ZipFile(stream, "a") as archive:
        if kind == "embedded":
            archive.writestr("word/embeddings/payload.bin", b"payload")
        else:
            archive.writestr(
                "word/_rels/custom.xml.rels",
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                f'<Relationship Id="rId100" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}" '
                'Target="https://example.com/resource" TargetMode="External"/></Relationships>',
            )
    client.force_authenticate(users[0])
    response = client.post(
        "/api/submissions/",
        {
            "title": "Relationships",
            "file": SimpleUploadedFile("links.docx", stream.getvalue()),
        },
        format="multipart",
    )
    assert response.status_code == expected
    if expected == 400:
        storage[0].assert_not_called()


def test_write_endpoints_require_csrf(users, pdf, storage, submission):
    client = APIClient(enforce_csrf_checks=True)
    client.force_login(users[0])
    assert (
        client.post(
            "/api/submissions/", {"title": "No CSRF", "file": pdf}, format="multipart"
        ).status_code
        == 403
    )
    client.force_login(users[2])
    assert (
        client.post(
            f"/api/submissions/{submission['id']}/decision/",
            {"status": "approved"},
            format="json",
        ).status_code
        == 403
    )
    assert Submission.objects.get(pk=submission["id"]).status == "pending"


@pytest.mark.parametrize("action", ["list", "detail", "file", "decision", "upload"])
def test_expired_session_errors_are_distinct_from_permission_failures(
    submission, action
):
    client = APIClient()
    base = f"/api/submissions/{submission['id']}/"
    if action in ("decision", "upload"):
        response = client.post(
            base + "decision/" if action == "decision" else "/api/submissions/", {}
        )
    else:
        response = client.get(
            {"list": "/api/submissions/", "detail": base, "file": base + "file/"}[
                action
            ]
        )
    assert response.status_code == 403
    assert response.data["code"] == "not_authenticated"


def test_forbidden_decision_does_not_report_session_expired(client, users, submission):
    client.force_authenticate(users[0])
    response = client.post(
        f"/api/submissions/{submission['id']}/decision/", {"status": "approved"}
    )
    assert response.status_code == 403
    assert response.data["code"] == "permission_denied"
